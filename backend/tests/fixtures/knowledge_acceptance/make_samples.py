"""Regenerate the 15-suffix acceptance samples (prep-time helper, not CI).

Requires: openpyxl, Pillow, python-docx, python-pptx, reportlab, xlwt.
Run from anywhere with a Python that has those:
    python make_samples.py <output-dir>

The two legacy Office formats are converted from dedicated sources
(``_legacy_doc.docx`` / ``_legacy_ppt.pptx``, also written by this script)
with Word / PowerPoint COM automation (see convert_legacy.ps1 next to this
file); on machines without Office, copying an OOXML file to ``.doc``/``.ppt``
is NOT acceptable — the acceptance requires real legacy binaries.

Every sample carries a unique marker token (``ACCEPT8-<KIND>-<digits>``) so a
parse/index/retrieval pass can verify that THIS file's content actually landed
in the knowledge base, rather than trusting the document name alone.
"""

from __future__ import annotations

import sys
from pathlib import Path

#: marker used per suffix; the acceptance question checks it end-to-end
MARKERS = {
    "md": "MD-3174",
    "markdown": "MARKDOWN-8620",
    "txt": "TXT-4509",
    "csv": "CSV-5527",
    "tsv": "TSV-6118",
    "xlsx": "XLSX-7302",
    "xls": "XLS-9266",
    "pdf": "PDF-1053",
    "docx": "DOCX-2287",
    "doc": "DOC-3390",
    "pptx": "PPTX-4416",
    "ppt": "PPT-5570",
    "png": "PNG-6633",
    "jpg": "JPG-7745",
    "jpeg": "JPEG-8812",
}


def _text(kind: str, label: str | None = None) -> str:
    label = label or kind
    return f"# 知识库验收样例（{label}）\n\n本文件是 15 后缀验收的 {label} 样例，用于验证上传、解析与索引路径。\n本文件的验收标记是 ACCEPT8-{MARKERS[kind]}，检索时以该标记为准。\n样例正文到此结束。\n"


def _font(size: int):
    from PIL import ImageFont

    for candidate in (r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\simhei.ttf"):
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default(size)


def _image(label: str, marker: str) -> object:
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (900, 500), "white")
    draw = ImageDraw.Draw(img)
    draw.rectangle((8, 8, 892, 492), outline="black", width=3)
    draw.text((40, 40), f"验收样例图（{label}）", fill="black", font=_font(36))
    draw.text((40, 140), "本图片的验收标记：", fill="black", font=_font(30))
    draw.text((40, 200), f"ACCEPT8-{marker}", fill="#b91c1c", font=_font(46))
    draw.text((40, 320), "用于验证图片上传、解析与检索路径。", fill="black", font=_font(26))
    return img


def main(out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)

    # plain text family -----------------------------------------------------
    (out_dir / "sample.md").write_text(_text("md", "Markdown"), encoding="utf-8")
    (out_dir / "sample.markdown").write_text(_text("markdown", "Markdown"), encoding="utf-8")
    (out_dir / "sample.txt").write_text(_text("txt", "纯文本"), encoding="utf-8")

    # delimited tables ------------------------------------------------------
    csv_lines = f"区域,产品,季度,销售额\n华东,样例甲,一季度,ACCEPT8-{MARKERS['csv']}\n华北,样例乙,二季度,88213\n"
    (out_dir / "sample.csv").write_text(csv_lines, encoding="utf-8")
    tsv_lines = f"区域\t产品\t季度\t销售额\n华南\t样例丙\t三季度\tACCEPT8-{MARKERS['tsv']}\n西北\t样例丁\t四季度\t93107\n"
    (out_dir / "sample.tsv").write_text(tsv_lines, encoding="utf-8")

    # workbooks -------------------------------------------------------------
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage

    wb = Workbook()
    ws = wb.active
    ws.title = "验收表"
    ws.append(["区域", "产品", "季度", "销售额"])
    ws.append(["东北", "样例戊", "一季度", f"ACCEPT8-{MARKERS['xlsx']}"])
    ws.append(["西南", "样例己", "二季度", 76420])
    ws.append(["下表图为本表的内嵌示意，图内标记与表格标记不同。", "", "", ""])
    diagram = out_dir / "_embedded.png"
    _image("XLSX 内嵌图", "IMG-8841").save(diagram)
    pic = XLImage(str(diagram))
    pic.anchor = "A5"
    ws.add_image(pic)
    wb.save(out_dir / "sample.xlsx")
    diagram.unlink()

    import xlwt

    book = xlwt.Workbook()
    sheet = book.add_sheet("验收表")
    for col, head in enumerate(["区域", "产品", "季度", "销售额"]):
        sheet.write(0, col, head)
    for col, value in enumerate(["中南", "样例庚", "三季度", f"ACCEPT8-{MARKERS['xls']}"]):
        sheet.write(1, col, value)
    sheet.write(2, 0, 61235)
    book.save(str(out_dir / "sample.xls"))

    # pdf -------------------------------------------------------------------
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    pdf = canvas.Canvas(str(out_dir / "sample.pdf"), pagesize=A4)
    pdf.setFont("STSong-Light", 18)
    pdf.drawString(80, 760, "知识库验收样例（PDF）")
    pdf.setFont("STSong-Light", 14)
    pdf.drawString(80, 710, "本文件的验收标记是 ACCEPT8-PDF-1053。")
    pdf.drawString(80, 680, "用于验证 PDF 上传、解析与索引路径。")
    pdf.showPage()
    pdf.save()

    # office documents ------------------------------------------------------
    from docx import Document

    doc = Document()
    doc.add_heading("知识库验收样例（DOCX）", level=1)
    doc.add_paragraph("本文件的验收标记是 ACCEPT8-DOCX-2287。")
    doc.add_paragraph("用于验证 Word 文档上传、解析与索引路径。")
    doc.save(str(out_dir / "sample.docx"))

    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "知识库验收样例（PPTX）"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
    frame = box.text_frame
    frame.text = "本文件的验收标记是 ACCEPT8-PPTX-4416。"
    frame.add_paragraph().text = "用于验证演示文稿上传、解析与索引路径。"
    prs.save(str(out_dir / "sample.pptx"))

    # legacy Office sources (sample.doc / sample.ppt come from convert_legacy.ps1)
    doc = Document()
    doc.add_heading("知识库验收样例（DOC）", level=1)
    doc.add_paragraph("本文件的验收标记是 ACCEPT8-DOC-3390。")
    doc.add_paragraph("用于验证旧版 Word（97-2003）上传、解析与索引路径。")
    doc.save(str(out_dir / "_legacy_doc.docx"))

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "知识库验收样例（PPT）"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
    frame = box.text_frame
    frame.text = "本文件的验收标记是 ACCEPT8-PPT-5570。"
    frame.add_paragraph().text = "用于验证旧版演示文稿（97-2003）上传、解析与索引路径。"
    prs.save(str(out_dir / "_legacy_ppt.pptx"))

    # images ----------------------------------------------------------------
    _image("PNG", MARKERS["png"]).save(out_dir / "sample.png")
    _image("JPG", MARKERS["jpg"]).save(out_dir / "sample.jpg", quality=92)
    _image("JPEG", MARKERS["jpeg"]).save(out_dir / "sample.jpeg", quality=92)

    print(f"wrote 13 samples + xlsx embedded image + legacy Office sources into {out_dir}")
    print("next: run convert_legacy.ps1 (Word/PowerPoint COM) to produce sample.doc / sample.ppt")


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent
    main(target)
